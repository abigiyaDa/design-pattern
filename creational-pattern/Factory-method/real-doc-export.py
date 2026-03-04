from abc import ABC, abstractmethod

# Install dependencies before running:
# ** pip install reportlab python-docx openpyxl**

# ------------------- Product Interface -------------------
class DocExporter(ABC):
    @abstractmethod
    def export(self, data, filename):
        pass

# ------------------- Concrete Products -------------------
# PDF
from reportlab.pdfgen import canvas

class PDFExporter(DocExporter):
    def export(self, data, filename):
        c = canvas.Canvas(filename)
        c.drawString(100, 750, data)
        c.save()
        return f"PDF saved as {filename}"

# Word
from docx import Document

class WordExporter(DocExporter):
    def export(self, data, filename):
        doc = Document()
        doc.add_paragraph(data)
        doc.save(filename)
        return f"Word file saved as {filename}"

# Excel
from openpyxl import Workbook

class ExcelExporter(DocExporter):
    def export(self, data, filename):
        wb = Workbook()
        ws = wb.active
        ws.append([data])  # Simple: write one row
        wb.save(filename)
        return f"Excel file saved as {filename}"

# ------------------- Factory -------------------
class DocExporterFactory(ABC):
    @abstractmethod
    def create_exporter(self) -> DocExporter:
        pass

# ------------------- Concrete factory -------------------

class PDFExporterFactory(DocExporterFactory):
    def create_exporter(self) -> DocExporter:
        return PDFExporter()

class WordExporterFactory(DocExporterFactory):
    def create_exporter(self) -> DocExporter:
        return WordExporter()

class ExcelExporterFactory(DocExporterFactory):
    def create_exporter(self) -> DocExporter:
        return ExcelExporter()

# ------------------- Client Code -------------------
def client_code(factory: DocExporterFactory, data, filename):
    exporter = factory.create_exporter()
    print(exporter.export(data, filename))

# ------------------- Test -------------------
if __name__ == "__main__":
    data = "This is a real document export example!"

    client_code(PDFExporterFactory(), data, "sample.pdf")
    client_code(WordExporterFactory(), data, "sample.docx")
    client_code(ExcelExporterFactory(), data, "sample.xlsx")

# Output:
# PDF saved as sample.pdf
# Word file saved as sample.docx
# Excel file saved as sample.xlsx